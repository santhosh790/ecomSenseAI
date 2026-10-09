import io
import sys
import types
import unittest
from unittest.mock import patch

import pandas as pd
from openpyxl import load_workbook

from application.reporting_service import export_excel, export_pdf, format_individual_order_summary


SAMPLE_ITEMS = pd.DataFrame(
    [{"Tamil Name": "வெங்காயம் (ONION)", "Total Quantity": 2, "Unit": "KG"}]
)


class IndividualDownloadMetadataTests(unittest.TestCase):
    def test_order_summary_counts_items_and_sums_kg_only(self):
        items = pd.DataFrame(
            [
                {"Total Quantity": 2, "Unit": "KG"},
                {"Total Quantity": 1.5, "Unit": "KGS"},
                {"Total Quantity": 4, "Unit": "EA"},
            ]
        )

        summary = format_individual_order_summary(items, "INR 1,250.00")

        self.assertEqual(
            summary,
            "Total Items: 3 | Total Weight: 3.5 KG | Total Amount: INR 1,250.00",
        )

    def test_excel_includes_order_metadata_but_defaults_do_not(self):
        summary = format_individual_order_summary(SAMPLE_ITEMS, "INR 1,250.00")
        content = export_excel(
            SAMPLE_ITEMS,
            client_name="MHS",
            footer_text=summary,
            purchase_order_number="PO-123",
            total_order_value="INR 1,250.00 (One Thousand Two Hundred Fifty Only)",
        )
        worksheet = load_workbook(io.BytesIO(content), data_only=True)["Vegetables"]
        self.assertIn("வாடிக்கையாளர்: MHS", worksheet["A3"].value)
        self.assertIn("PO Number: PO-123", worksheet["A3"].value)
        self.assertIsNone(worksheet["A4"].value)
        self.assertIsNone(worksheet["A5"].value)
        self.assertIn("Total Items: 1 | Total Weight: 2 KG | Total Amount: INR 1,250.00", worksheet["A10"].value)

        default_content = export_excel(SAMPLE_ITEMS)
        default_worksheet = load_workbook(io.BytesIO(default_content), data_only=True)["Vegetables"]
        self.assertIsNone(default_worksheet["A4"].value)
        self.assertIsNone(default_worksheet["A5"].value)

    def test_pdf_includes_optional_order_metadata(self):
        captured_html = []

        class FakeHTML:
            def __init__(self, string):
                captured_html.append(string)

            def write_pdf(self, output):
                output.write(b"%PDF-test")

        fake_weasyprint = types.ModuleType("weasyprint")
        fake_weasyprint.HTML = FakeHTML
        with patch.dict(sys.modules, {"weasyprint": fake_weasyprint}):
            export_pdf(
                SAMPLE_ITEMS,
                client_name="MHS",
                purchase_order_number="PO-123",
                total_order_value="INR 1,250.00",
                order_totals_summary="Total Items: 1 | Total Weight: 2 KG | Total Amount: INR 1,250.00",
            )

        self.assertIn("வாடிக்கையாளர்:</b> MHS    |    <b>PO Number:</b> PO-123", captured_html[0])
        self.assertNotIn("Total Order Value:</b>", captured_html[0])
        self.assertIn("Total Items: 1 | Total Weight: 2 KG | Total Amount: INR 1,250.00", captured_html[0])


if __name__ == "__main__":
    unittest.main()
